"""Document intake, text extraction and deterministic pre-checks for correspondence."""
from pathlib import Path
import re
import sqlite3
from datetime import datetime
from werkzeug.utils import secure_filename
from config import BASE_DIR, DATABASE_PATH

UPLOAD_DIR = BASE_DIR / "data" / "checker_uploads"
ALLOWED_EXTENSIONS = {".docx", ".pdf", ".txt"}

CHECKS = [
    ("nomor", "Nomor naskah/surat tercantum"),
    ("tanggal", "Tanggal naskah tercantum"),
    ("tujuan", "Tujuan/alamat penerima jelas"),
    ("perihal", "Perihal/subjek tercantum"),
    ("isi", "Isi naskah memuat maksud dan tindak lanjut yang jelas"),
    ("lampiran", "Lampiran disebutkan atau dinyatakan tidak ada"),
    ("penandatangan", "Nama/jabatan penandatangan tercantum"),
    ("bahasa", "Bahasa, ejaan, dan tanda baca telah ditinjau manual"),
    ("format", "Format sesuai tata naskah dinas yang berlaku"),
]


def _connect():
    conn = sqlite3.connect(DATABASE_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_checker_table():
    with _connect() as conn:
        conn.execute("""CREATE TABLE IF NOT EXISTS correspondence_checks (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            document_type TEXT NOT NULL,
            title TEXT NOT NULL,
            document_number TEXT DEFAULT '',
            document_date TEXT DEFAULT '',
            sender TEXT DEFAULT '',
            recipient TEXT DEFAULT '',
            subject TEXT DEFAULT '',
            attachment_name TEXT NOT NULL,
            stored_path TEXT NOT NULL,
            extracted_text TEXT DEFAULT '',
            automated_findings TEXT DEFAULT '[]',
            source_findings TEXT DEFAULT '[]',
            checklist_json TEXT DEFAULT '{}',
            status TEXT NOT NULL DEFAULT 'Menunggu Pemeriksaan',
            checker_notes TEXT DEFAULT '',
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        )""")
        # Backward-compatible migration for databases created by earlier checker versions.
        columns = {row[1] for row in conn.execute("PRAGMA table_info(correspondence_checks)").fetchall()}
        if "source_findings" not in columns:
            conn.execute("ALTER TABLE correspondence_checks ADD COLUMN source_findings TEXT DEFAULT '[]'")


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".txt":
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix == ".docx":
        from docx import Document
        doc = Document(str(path))
        chunks = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                chunks.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(chunks)
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise ValueError("Pembacaan PDF memerlukan paket pypdf. Jalankan: pip install pypdf") from exc
        reader = PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    raise ValueError("Format file belum didukung. Gunakan DOCX, PDF, atau TXT.")


def precheck(text: str):
    normalized = re.sub(r"\s+", " ", text).strip()
    lower = normalized.lower()
    findings = []
    if not normalized:
        findings.append({"severity": "warning", "code": "empty_text", "message": "Teks tidak dapat diekstrak. PDF mungkin hasil scan; lakukan pemeriksaan manual/OCR."})
        return findings
    patterns = [
        (r"\bnomor\b|\bno\.?\s*[:/]", "Nomor naskah terdeteksi", "Nomor surat tidak terdeteksi jelas; verifikasi secara manual."),
        (r"\bperihal\b|\bhal\s*:", "Perihal terdeteksi", "Perihal tidak terdeteksi; pastikan sesuai jenis naskah."),
        (r"\blampiran\b", "Bagian lampiran terdeteksi", "Kata/bagian lampiran tidak terdeteksi; pastikan lampiran ada atau dinyatakan nihil."),
        (r"\bkepada\b|\byth\.?\b|\byang terhormat\b", "Tujuan/penerima terdeteksi", "Tujuan/penerima tidak terdeteksi jelas."),
        (r"\b(tanggal|jakarta|tangerang|bandung)\b", "Indikasi tanggal/tempat terdeteksi", "Tanggal/tempat naskah perlu diverifikasi."),
    ]
    for pattern, success, warning in patterns:
        if re.search(pattern, lower):
            findings.append({"severity": "info", "code": "pattern_found", "message": success})
        else:
            findings.append({"severity": "warning", "code": "pattern_missing", "message": warning})
    suspicious = re.findall(r"\b(?:yg|dgn|utk|krn|tdk|sdh|aja|gak)\b", lower)
    if suspicious:
        findings.append({"severity": "warning", "code": "informal_words", "message": "Ditemukan kemungkinan singkatan informal: " + ", ".join(sorted(set(suspicious))[:8]) + ". Periksa konteks dan gunakan bahasa kedinasan bila sesuai."})
    if len(normalized) < 100:
        findings.append({"severity": "warning", "code": "short_text", "message": "Teks hasil ekstraksi cukup singkat; pastikan seluruh halaman dan isi terbaca."})
    findings.append({"severity": "info", "code": "human_review", "message": "Pemeriksaan otomatis hanya skrining awal. Ketepatan nomor, dasar hukum, substansi, tanda tangan, lampiran fisik, dan kesesuaian tata naskah harus diverifikasi checker."})
    return findings


def list_checks():
    with _connect() as conn:
        rows = conn.execute("""SELECT id, document_type, title, document_number, document_date,
            sender, recipient, subject, attachment_name, status, checker_notes, created_at, updated_at
            FROM correspondence_checks ORDER BY id DESC""").fetchall()
        return [dict(row) for row in rows]


def create_check(form, uploaded):
    original = secure_filename(uploaded.filename or "")
    suffix = Path(original).suffix.lower()
    if not original or suffix not in ALLOWED_EXTENSIONS:
        raise ValueError("File wajib berupa DOCX, PDF, atau TXT.")
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    stored_name = f"{datetime.now().strftime('%Y%m%d%H%M%S%f')}_{original}"
    path = UPLOAD_DIR / stored_name
    uploaded.save(path)
    try:
        text = extract_text(path)
        findings = precheck(text)
        from services.analysis_provider import get_analysis_provider
        source_findings = get_analysis_provider().analyze((form.get("document_type") or "Nota Dinas").strip(), text)
        now = datetime.now().isoformat(timespec="seconds")
        doc_type = (form.get("document_type") or "Nota Dinas").strip()
        title = (form.get("title") or Path(original).stem).strip()[:200]
        with _connect() as conn:
            cur = conn.execute("""INSERT INTO correspondence_checks
                (document_type,title,document_number,document_date,sender,recipient,subject,attachment_name,stored_path,
                 extracted_text,automated_findings,source_findings,checklist_json,status,checker_notes,created_at,updated_at)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (doc_type, title, form.get("document_number", "").strip(), form.get("document_date", "").strip(),
                 form.get("sender", "").strip(), form.get("recipient", "").strip(), form.get("subject", "").strip(),
                 original, str(path), text[:200000], __import__('json').dumps(findings, ensure_ascii=False),
                 __import__('json').dumps(source_findings, ensure_ascii=False), '{}',
                 "Menunggu Pemeriksaan", "", now, now))
            check_id = cur.lastrowid
            conn.execute("INSERT INTO checker_audit_log(entity_type,entity_id,action,detail_json,created_at) VALUES(?,?,?,?,?)",
                ("correspondence_check", check_id, "document_checked_against_rules",
                 __import__('json').dumps({"source_findings":len(source_findings),"rule_ids":[f.get("rule_id") for f in source_findings]},ensure_ascii=False), now))
            return get_check(check_id)
    except Exception:
        path.unlink(missing_ok=True)
        raise


def get_check(check_id):
    import json
    with _connect() as conn:
        row = conn.execute("SELECT * FROM correspondence_checks WHERE id=?", (check_id,)).fetchone()
        if not row:
            return None
        result = dict(row)
        result["automated_findings"] = json.loads(result["automated_findings"] or "[]")
        result["source_findings"] = json.loads(result.get("source_findings") or "[]")
        result["checklist"] = json.loads(result.pop("checklist_json") or "{}")
        result.pop("stored_path", None)
        return result


def update_check(check_id, data):
    import json
    allowed_statuses = {"Menunggu Pemeriksaan", "Sedang Diperiksa", "Perlu Revisi", "Siap Diteruskan", "Selesai"}
    status = data.get("status", "Menunggu Pemeriksaan")
    if status not in allowed_statuses:
        raise ValueError("Status pemeriksaan tidak valid.")
    checklist = data.get("checklist", {})
    if not isinstance(checklist, dict):
        raise ValueError("Checklist tidak valid.")
    now = datetime.now().isoformat(timespec="seconds")
    with _connect() as conn:
        cur = conn.execute("UPDATE correspondence_checks SET status=?, checker_notes=?, checklist_json=?, updated_at=? WHERE id=?",
            (status, str(data.get("checker_notes", ""))[:10000], json.dumps(checklist, ensure_ascii=False), now, check_id))
        if not cur.rowcount:
            return None
        conn.execute("INSERT INTO checker_audit_log(entity_type,entity_id,action,detail_json,created_at) VALUES(?,?,?,?,?)",
            ("correspondence_check", check_id, "review_updated", json.dumps({"status":status,"checklist":checklist},ensure_ascii=False), now))
    return get_check(check_id)


def get_stored_file(check_id):
    with _connect() as conn:
        row = conn.execute("SELECT stored_path, attachment_name FROM correspondence_checks WHERE id=?", (check_id,)).fetchone()
        if not row:
            return None
        return Path(row["stored_path"]), row["attachment_name"]
