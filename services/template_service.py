from docx import Document


def load_template(path):

    return Document(path)


def replace_placeholders(doc, replacements):

    for paragraph in doc.paragraphs:

        for key, value in replacements.items():

            if key in paragraph.text:

                paragraph.text = paragraph.text.replace(
                    key,
                    value
                )